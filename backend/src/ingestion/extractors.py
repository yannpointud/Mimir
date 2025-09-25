"""
Text extraction from various document formats
Provides extractors for PDF, DOCX, TXT, and Markdown files
"""

from abc import ABC, abstractmethod
from pathlib import Path
from typing import Optional, Dict, Any
import logging

# Import document processing libraries
try:
    import pypdf
except ImportError:
    pypdf = None

try:
    from docx import Document as DocxDocument
except ImportError:
    DocxDocument = None

try:
    import markdown
except ImportError:
    markdown = None

from ..monitoring.logger import setup_logger

logger = setup_logger(__name__)

class ExtractionError(Exception):
    """Base exception for text extraction errors"""
    pass

class UnsupportedFormatError(ExtractionError):
    """Raised when file format is not supported"""
    pass

class TextExtractor(ABC):
    """Abstract base class for text extractors"""

    @abstractmethod
    def extract(self, file_path: str) -> str:
        """
        Extract text from file

        Args:
            file_path: Path to the file

        Returns:
            Extracted text content

        Raises:
            ExtractionError: If extraction fails
        """
        pass

    @abstractmethod
    def can_extract(self, file_path: str) -> bool:
        """
        Check if this extractor can handle the file

        Args:
            file_path: Path to the file

        Returns:
            True if extractor can handle this file type
        """
        pass

    def get_metadata(self, file_path: str) -> Dict[str, Any]:
        """
        Extract metadata from file (optional override)

        Args:
            file_path: Path to the file

        Returns:
            Dictionary with metadata
        """
        file_path_obj = Path(file_path)
        stat = file_path_obj.stat()

        return {
            'file_size': stat.st_size,
            'modified_time': stat.st_mtime,
            'file_extension': file_path_obj.suffix.lower(),
            'file_name': file_path_obj.name
        }

class PDFExtractor(TextExtractor):
    """Extract text from PDF files using pypdf"""

    def __init__(self):
        if pypdf is None:
            raise ImportError("pypdf is required for PDF extraction. Install with: pip install pypdf")

    def can_extract(self, file_path: str) -> bool:
        """Check if file is PDF"""
        return Path(file_path).suffix.lower() == '.pdf'

    def extract(self, file_path: str) -> str:
        """Extract text from PDF file"""
        if not self.can_extract(file_path):
            raise UnsupportedFormatError(f"PDF extractor cannot handle file: {file_path}")

        try:
            text_parts = []

            with open(file_path, 'rb') as file:
                reader = pypdf.PdfReader(file)

                # Check if PDF is encrypted
                if reader.is_encrypted:
                    logger.warning(f"PDF is encrypted, attempting to decrypt: {file_path}")
                    try:
                        reader.decrypt('')  # Try empty password
                    except Exception:
                        raise ExtractionError(f"Cannot decrypt encrypted PDF: {file_path}")

                # Extract text from each page
                for page_num, page in enumerate(reader.pages):
                    try:
                        page_text = page.extract_text()
                        if page_text.strip():  # Only add non-empty pages
                            text_parts.append(page_text)
                    except Exception as e:
                        logger.warning(f"Failed to extract text from page {page_num + 1} in {file_path}: {str(e)}")
                        continue

            if not text_parts:
                logger.warning(f"No text extracted from PDF: {file_path}")
                return ""

            extracted_text = '\n\n'.join(text_parts)

            logger.info(
                f"Extracted {len(extracted_text)} characters from PDF: {Path(file_path).name}",
                extra={'extra_fields': {
                    'file_path': file_path,
                    'pages': len(reader.pages),
                    'characters': len(extracted_text),
                    'type': 'pdf_extraction'
                }}
            )

            return extracted_text

        except Exception as e:
            error_msg = f"Failed to extract text from PDF {file_path}: {str(e)}"
            logger.error(error_msg)
            raise ExtractionError(error_msg)

class DOCXExtractor(TextExtractor):
    """Extract text from DOCX files using python-docx"""

    def __init__(self):
        if DocxDocument is None:
            raise ImportError("python-docx is required for DOCX extraction. Install with: pip install python-docx")

    def can_extract(self, file_path: str) -> bool:
        """Check if file is DOCX"""
        return Path(file_path).suffix.lower() == '.docx'

    def extract(self, file_path: str) -> str:
        """Extract text from DOCX file"""
        if not self.can_extract(file_path):
            raise UnsupportedFormatError(f"DOCX extractor cannot handle file: {file_path}")

        try:
            doc = DocxDocument(file_path)
            text_parts = []

            # Extract text from paragraphs
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text_parts.append(paragraph.text)

            # Extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = []
                    for cell in row.cells:
                        if cell.text.strip():
                            row_text.append(cell.text.strip())
                    if row_text:
                        text_parts.append(' | '.join(row_text))

            if not text_parts:
                logger.warning(f"No text extracted from DOCX: {file_path}")
                return ""

            extracted_text = '\n\n'.join(text_parts)

            logger.info(
                f"Extracted {len(extracted_text)} characters from DOCX: {Path(file_path).name}",
                extra={'extra_fields': {
                    'file_path': file_path,
                    'paragraphs': len(doc.paragraphs),
                    'tables': len(doc.tables),
                    'characters': len(extracted_text),
                    'type': 'docx_extraction'
                }}
            )

            return extracted_text

        except Exception as e:
            error_msg = f"Failed to extract text from DOCX {file_path}: {str(e)}"
            logger.error(error_msg)
            raise ExtractionError(error_msg)

class TXTExtractor(TextExtractor):
    """Extract text from plain text files"""

    SUPPORTED_ENCODINGS = ['utf-8', 'utf-16', 'latin-1', 'cp1252']

    def can_extract(self, file_path: str) -> bool:
        """Check if file is text file"""
        ext = Path(file_path).suffix.lower()
        return ext in ['.txt', '.text', '.log', '.csv', '.json', '.xml']

    def extract(self, file_path: str) -> str:
        """Extract text from text file with encoding detection"""
        if not self.can_extract(file_path):
            raise UnsupportedFormatError(f"TXT extractor cannot handle file: {file_path}")

        # Try different encodings
        for encoding in self.SUPPORTED_ENCODINGS:
            try:
                with open(file_path, 'r', encoding=encoding) as file:
                    content = file.read()

                logger.info(
                    f"Extracted {len(content)} characters from TXT: {Path(file_path).name}",
                    extra={'extra_fields': {
                        'file_path': file_path,
                        'encoding': encoding,
                        'characters': len(content),
                        'type': 'txt_extraction'
                    }}
                )

                return content

            except UnicodeDecodeError:
                continue
            except Exception as e:
                error_msg = f"Failed to read text file {file_path} with encoding {encoding}: {str(e)}"
                logger.warning(error_msg)
                continue

        # If all encodings failed
        error_msg = f"Failed to extract text from {file_path}: Unable to decode with any supported encoding"
        logger.error(error_msg)
        raise ExtractionError(error_msg)

class MarkdownExtractor(TextExtractor):
    """Extract text from Markdown files"""

    def __init__(self, convert_to_html: bool = False):
        """
        Initialize Markdown extractor

        Args:
            convert_to_html: If True, convert markdown to HTML first (requires markdown library)
        """
        self.convert_to_html = convert_to_html
        if convert_to_html and markdown is None:
            raise ImportError("markdown is required for HTML conversion. Install with: pip install markdown")

    def can_extract(self, file_path: str) -> bool:
        """Check if file is Markdown"""
        ext = Path(file_path).suffix.lower()
        return ext in ['.md', '.markdown', '.mdown', '.mkd']

    def extract(self, file_path: str) -> str:
        """Extract text from Markdown file"""
        if not self.can_extract(file_path):
            raise UnsupportedFormatError(f"Markdown extractor cannot handle file: {file_path}")

        try:
            # First, read the raw markdown
            with open(file_path, 'r', encoding='utf-8') as file:
                content = file.read()

            if self.convert_to_html and markdown is not None:
                # Convert markdown to HTML (preserves structure better)
                content = markdown.markdown(content)

            logger.info(
                f"Extracted {len(content)} characters from Markdown: {Path(file_path).name}",
                extra={'extra_fields': {
                    'file_path': file_path,
                    'characters': len(content),
                    'converted_to_html': self.convert_to_html,
                    'type': 'markdown_extraction'
                }}
            )

            return content

        except Exception as e:
            error_msg = f"Failed to extract text from Markdown {file_path}: {str(e)}"
            logger.error(error_msg)
            raise ExtractionError(error_msg)

class ExtractorFactory:
    """Factory for creating appropriate text extractors"""

    def __init__(self):
        self._extractors = []
        self._initialize_extractors()

    def _initialize_extractors(self):
        """Initialize available extractors"""
        # Add extractors in order of preference

        # PDF extractor
        try:
            self._extractors.append(PDFExtractor())
            logger.debug("PDF extractor initialized")
        except ImportError as e:
            logger.warning(f"PDF extractor not available: {e}")

        # DOCX extractor
        try:
            self._extractors.append(DOCXExtractor())
            logger.debug("DOCX extractor initialized")
        except ImportError as e:
            logger.warning(f"DOCX extractor not available: {e}")

        # Markdown extractor (with HTML conversion if available)
        try:
            self._extractors.append(MarkdownExtractor(convert_to_html=True))
            logger.debug("Markdown extractor initialized (with HTML conversion)")
        except ImportError:
            self._extractors.append(MarkdownExtractor(convert_to_html=False))
            logger.debug("Markdown extractor initialized (plain text)")

        # Plain text extractor (always available)
        self._extractors.append(TXTExtractor())
        logger.debug("TXT extractor initialized")

        logger.info(f"Initialized {len(self._extractors)} text extractors")

    def get_extractor(self, file_path: str) -> Optional[TextExtractor]:
        """
        Get appropriate extractor for file

        Args:
            file_path: Path to the file

        Returns:
            TextExtractor instance or None if no suitable extractor found
        """
        for extractor in self._extractors:
            if extractor.can_extract(file_path):
                return extractor

        logger.warning(f"No suitable extractor found for file: {file_path}")
        return None

    def extract_text(self, file_path: str) -> str:
        """
        Extract text using appropriate extractor

        Args:
            file_path: Path to the file

        Returns:
            Extracted text

        Raises:
            UnsupportedFormatError: If no suitable extractor found
            ExtractionError: If extraction fails
        """
        extractor = self.get_extractor(file_path)
        if not extractor:
            raise UnsupportedFormatError(f"No extractor available for file: {file_path}")

        return extractor.extract(file_path)

    def get_supported_extensions(self) -> list[str]:
        """Get list of supported file extensions"""
        extensions = set()

        # Test each extractor with common extensions
        test_extensions = ['.pdf', '.docx', '.txt', '.md', '.markdown', '.log', '.csv', '.json', '.xml']

        for ext in test_extensions:
            test_file = f"test{ext}"
            for extractor in self._extractors:
                if extractor.can_extract(test_file):
                    extensions.add(ext)
                    break

        return sorted(list(extensions))